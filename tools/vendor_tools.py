import os
import csv
import openpyxl
from typing import List, Dict, Any, Tuple
from core.tool_registry import register_tool


def normalize_column_name(col: str) -> str:
    """Normalizes header variants into canonical field names."""
    c = col.lower().replace("_", " ").replace("-", " ").strip()
    if any(k in c for k in ["vendor sku", "sku", "item code", "product id", "part number"]):
        return "sku"
    elif any(k in c for k in ["product name", "item name", "name", "title", "description"]):
        return "product_name"
    elif any(k in c for k in ["price", "cost", "wholesale"]):
        return "price"
    elif any(k in c for k in ["category", "dept"]):
        return "category"
    elif any(k in c for k in ["stock", "qty", "quantity"]):
        return "stock_qty"
    elif any(k in c for k in ["lead time", "lead"]):
        return "lead_time_days"
    return c.replace(" ", "_")


@register_tool(name="vendor_spreadsheet_parser", description="Parses Excel (xlsx) or CSV vendor product catalogs", category="vendor")
def parse_vendor_file(filepath: str = "data/vendor_products_dirty.xlsx") -> Tuple[List[str], List[Dict[str, Any]]]:
    if not os.path.exists(filepath):
        # Fallback to CSV if xlsx not found
        csv_alt = filepath.replace(".xlsx", ".csv")
        if os.path.exists(csv_alt):
            filepath = csv_alt
        else:
            raise FileNotFoundError(f"Vendor spreadsheet file not found: {filepath}")

    raw_rows = []
    headers = []

    if filepath.endswith(".xlsx"):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        sheet = wb.active
        headers = [str(sheet.cell(1, c).value or "").strip() for c in range(1, sheet.max_column + 1)]
        for r in range(2, sheet.max_row + 1):
            row_dict = {}
            has_data = False
            for c in range(1, sheet.max_column + 1):
                val = sheet.cell(r, c).value
                if val is not None:
                    has_data = True
                row_dict[headers[c - 1]] = val
            if has_data:
                raw_rows.append(row_dict)
    else:
        with open(filepath, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            for row in reader:
                raw_rows.append(dict(row))

    return headers, raw_rows


@register_tool(name="vendor_data_cleaner", description="Normalizes columns, validates required SKU/Name, flags invalid rows, and cleans dataset", category="vendor")
def process_and_clean_vendor_data(filepath: str = "data/vendor_products_dirty.xlsx") -> Dict[str, Any]:
    raw_headers, raw_rows = parse_vendor_file(filepath)

    # Column normalization mapping
    normalized_header_map = {orig: normalize_column_name(orig) for orig in raw_headers}

    valid_rows = []
    invalid_rows = []

    for idx, row in enumerate(raw_rows, start=2):  # row 2 is first data row
        normalized_row = {}
        for orig_key, val in row.items():
            norm_key = normalized_header_map.get(orig_key, orig_key)
            normalized_row[norm_key] = val

        sku_val = str(normalized_row.get("sku") or "").strip()
        name_val = str(normalized_row.get("product_name") or "").strip()

        # Decision Logic: Rows missing SKU or product name are invalid
        invalid_reasons = []
        if not sku_val or sku_val.lower() in ["none", "null", "nan", ""]:
            invalid_reasons.append("Missing required field: SKU")
        if not name_val or name_val.lower() in ["none", "null", "nan", ""]:
            invalid_reasons.append("Missing required field: Product Name")

        # Additional format validation
        price_val = normalized_row.get("price")
        if price_val is not None:
            try:
                price_float = float(price_val)
                if price_float < 0:
                    invalid_reasons.append(f"Invalid negative wholesale price: {price_float}")
                normalized_row["price"] = price_float
            except (ValueError, TypeError):
                invalid_reasons.append(f"Non-numeric price format: '{price_val}'")

        if invalid_reasons:
            invalid_rows.append({
                "row_number": idx,
                "raw_data": row,
                "reasons": invalid_reasons
            })
        else:
            valid_rows.append({
                "row_number": idx,
                "cleaned_data": normalized_row
            })

    # Save cleaned dataset to output file
    os.makedirs("data/output", exist_ok=True)
    cleaned_file_path = "data/output/cleaned_vendor_dataset.csv"
    if valid_rows:
        fieldnames = list(valid_rows[0]["cleaned_data"].keys())
        with open(cleaned_file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for r in valid_rows:
                writer.writerow(r["cleaned_data"])

    return {
        "file_processed": os.path.basename(filepath),
        "total_rows_inspected": len(raw_rows),
        "valid_rows_count": len(valid_rows),
        "invalid_rows_count": len(invalid_rows),
        "cleanliness_rate_pct": round((len(valid_rows) / len(raw_rows)) * 100.0, 1) if raw_rows else 0.0,
        "cleaned_file_saved_at": cleaned_file_path,
        "valid_sample": [v["cleaned_data"] for v in valid_rows[:5]],
        "invalid_row_report": invalid_rows
    }
