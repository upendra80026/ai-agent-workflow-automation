import openpyxl
from typing import Dict, List, Optional
import os
import re
from core.models import WorkflowDefinition


def parse_steps(steps_text: str) -> List[str]:
    """Parse steps separated by arrow (→ or ->) or commas."""
    if not steps_text:
        return []
    # Replace unicode arrow or ascii arrow
    cleaned = steps_text.replace("→", "->")
    if "->" in cleaned:
        raw_steps = cleaned.split("->")
    else:
        raw_steps = cleaned.split(",")
    return [s.strip() for s in raw_steps if s.strip()]


def parse_tools(tools_text: str) -> List[str]:
    """Parse tools separated by semicolon or comma."""
    if not tools_text:
        return []
    parts = re.split(r"[;,]", tools_text)
    return [p.strip() for p in parts if p.strip()]


def load_workflows_from_excel(filepath: str = "AI_Agent_Workflow_Assessment.xlsx") -> Dict[str, WorkflowDefinition]:
    """
    Dynamically loads workflow definitions directly from the Excel file.
    Does NOT hard-code workflows. Supports any number of workflows added to the sheet.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Workflow Excel file not found at: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    if "Workflows" not in wb.sheetnames:
        raise ValueError(f"'Workflows' sheet not found in {filepath}. Available: {wb.sheetnames}")

    sheet = wb["Workflows"]
    headers = [str(sheet.cell(1, c).value or "").strip() for c in range(1, sheet.max_column + 1)]

    # Map headers
    header_indices = {h: idx + 1 for idx, h in enumerate(headers)}

    workflows: Dict[str, WorkflowDefinition] = {}

    for row_idx in range(2, sheet.max_row + 1):
        wf_id = sheet.cell(row_idx, header_indices.get("Workflow_ID", 1)).value
        if not wf_id:
            continue
        wf_id = str(wf_id).strip()

        wf_name = str(sheet.cell(row_idx, header_indices.get("Workflow_Name", 2)).value or "").strip()
        trigger = str(sheet.cell(row_idx, header_indices.get("Trigger", 3)).value or "").strip()
        inputs = str(sheet.cell(row_idx, header_indices.get("Inputs", 4)).value or "").strip()
        steps_raw = str(sheet.cell(row_idx, header_indices.get("Steps", 5)).value or "").strip()
        decision_logic = str(sheet.cell(row_idx, header_indices.get("Decision_Logic", 6)).value or "").strip()
        tools_raw = str(sheet.cell(row_idx, header_indices.get("Tools_Required", 7)).value or "").strip()
        expected_output = str(sheet.cell(row_idx, header_indices.get("Expected_Output", 8)).value or "").strip()

        parsed_steps = parse_steps(steps_raw)
        parsed_tools = parse_tools(tools_raw)

        wf_def = WorkflowDefinition(
            workflow_id=wf_id,
            workflow_name=wf_name,
            trigger=trigger,
            inputs=inputs,
            steps_raw=steps_raw,
            steps=parsed_steps,
            decision_logic=decision_logic,
            tools_required=parsed_tools,
            expected_output=expected_output
        )
        workflows[wf_id] = wf_def

    return workflows


def load_test_questions_from_excel(filepath: str = "AI_Agent_Workflow_Assessment.xlsx") -> List[Dict[str, str]]:
    """Loads evaluation test questions directly from the Test_Questions sheet."""
    if not os.path.exists(filepath):
        return []

    wb = openpyxl.load_workbook(filepath, data_only=True)
    if "Test_Questions" not in wb.sheetnames:
        return []

    sheet = wb["Test_Questions"]
    questions = []
    for row_idx in range(2, sheet.max_row + 1):
        wf_id = sheet.cell(row_idx, 1).value
        test_request = sheet.cell(row_idx, 2).value
        what_to_check = sheet.cell(row_idx, 3).value
        if wf_id and test_request:
            questions.append({
                "workflow_id": str(wf_id).strip(),
                "test_request": str(test_request).strip(),
                "what_to_check": str(what_to_check or "").strip()
            })
    return questions
